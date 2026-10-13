import sys, os, io
sys.path.insert(0, os.path.abspath('.'))
import openpyxl
from cloud.app_cloud import generate_invoice_excel

test_invoice = {
    'invoice_number': '62/2026-27',
    'invoice_date': '2026-10-14',
    'buyer_name': 'Test Six Items Metal Works',
    'buyer_details': ['Test Six Items Metal Works', 'Main Market', 'Patna - 800001', 'STATE- BIHAR CODE-10', 'GSTIN - 10AABCT1234F1Z5'],
    'transport_mode': 'By Roadways',
    'delivery_charge': 600.0,
    'tax_type': 'IGST',
    'items': [
        {'description': 'Aluminium Utensils (5 Bags)', 'quantity': 150.5, 'rate': 370.0, 'hsn': '76151030'},
        {'description': 'Aluminium Circles (3 Bags)', 'quantity': 90.0, 'rate': 365.0, 'hsn': '76151030'},
        {'description': 'Aluminium Sheets (4 Bags)', 'quantity': 120.25, 'rate': 380.0, 'hsn': '76151030'},
        {'description': 'Aluminium Wire (2 Bags)', 'quantity': 60.0, 'rate': 390.0, 'hsn': '76151030'},
        {'description': 'Aluminium Scrap (1 Bag)', 'quantity': 45.75, 'rate': 290.0, 'hsn': '76151030'},
        {'description': 'Aluminium Topi (6 Bags)', 'quantity': 210.0, 'rate': 375.0, 'hsn': '76151030'}
    ]
}

excel_bytes = generate_invoice_excel(test_invoice)
wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
ws = wb.active

print('=== 6 ITEMS CHECK ===')
for idx in range(6):
    r = 22 + idx
    desc = ws[f'A{r}'].value
    qty = ws[f'F{r}'].value
    rate = ws[f'G{r}'].value
    hsn = ws[f'H{r}'].value
    amt = ws[f'I{r}'].value
    hsn_bold = ws[f'H{r}'].font.bold
    print(f'Item {idx+1} (Row {r}): desc={desc!r}, qty={qty}, rate={rate}, hsn={hsn} (bold={hsn_bold}), amt={amt}')

print('\n=== EMPTY ROWS (28-31) CHECK ===')
for r in range(28, 32):
    t_border = ws[f'A{r}'].border.top.style if ws[f'A{r}'].border and ws[f'A{r}'].border.top else None
    print(f'Row {r}: A={ws[f"A{r}"].value!r}, top_border={t_border}')

print('\n=== TOTALS & TAX (32-38) CHECK ===')
for r in range(32, 39):
    print(f'Row {r}: C={ws[f"C{r}"].value!r}, E={ws[f"E{r}"].value!r}, I={ws[f"I{r}"].value!r}, hidden={ws.row_dimensions[r].hidden}')
